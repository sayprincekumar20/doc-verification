"""RGF Customer Information Sheet (.xlsx), other labelled items, contact proposals, and the
fixes found on a real customer (business name misread, Manila e-permit validity)."""

import hashlib
import io
import zipfile
from datetime import date

import openpyxl

from app.assessment.assess import DocInput, assess
from app.assessment.matching import LIKELY_SAME, match_business
from app.extraction.spreadsheet import (
    extract_customer_info_sheet,
    is_customer_info_sheet,
    label_value_pairs,
)
from app.extraction.validate import check_validity
from app.pipeline.file_checks import check_file
from tests.test_assessment import SNAPSHOT, bir, dti, permit

TODAY = date(2026, 10, 8)


def cis_bytes(**overrides) -> bytes:
    wb = openpyxl.Workbook()
    ws = wb.active
    ws["A1"] = "RARE GLOBAL FOOD TRADING CORPORATION"
    ws["A2"] = "CUSTOMER INFORMATION SHEET  •  Confidential"
    rows = {
        5: ("Date Submitted:", " October 8, 2026", "Customer Code (RG use only):", None),
        8: ("Company Name (Legal Entity) *", " Juan's Meat Shop",
            "Business Style / Trade Name *", " Meat Retail"),
        10: ("TIN (incl. branch code) *", overrides.get("tin", " 123-456-789-00000"),
             "VAT Type *", None),
        15: ("Billing Address *", "PUROK 1 MASAGANA 4033 BAY", "Delivery Address (if different)",
             None),
        33: ("Primary Contact Person *", " Juan Dela Cruz", "Position / Department *", " Owner"),
        34: ("Email Address *", " juan@example.com", "Mobile / Landline (multiple) *", 9171234567),
        35: ("Backup Contact Person", " Maria", "Backup Email / Number", " maria@example.com"),
    }
    for r, (b, c, e, fv) in rows.items():
        ws[f"B{r}"], ws[f"C{r}"], ws[f"E{r}"], ws[f"F{r}"] = b, c, e, fv
    for col, text in zip("BCDEF", ["Reference Type", "Company / Bank Name", "Contact Person",
                                   "Contact No. / Email", "Relationship / Years"], strict=True):
        ws[f"{col}51"] = text   # table header: must not become label/value pairs
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def as_plain_zip(xlsx: bytes) -> bytes:
    """Rewrite without [Content_Types].xml first, like files that file-type sniffing calls zip."""
    src = zipfile.ZipFile(io.BytesIO(xlsx))
    out = io.BytesIO()
    with zipfile.ZipFile(out, "w") as z:
        z.writestr("_rels/", b"")
        for name in src.namelist():
            z.writestr(name, src.read(name))
    return out.getvalue()


def test_sheet_is_read_from_cells():
    data = cis_bytes()
    assert is_customer_info_sheet(data)
    r = extract_customer_info_sheet(data)
    fields = r["fields"]
    assert fields["tin"]["value"] == "123-456-789-00000"
    assert fields["company_name"]["value"] == "Juan's Meat Shop"
    assert fields["phone"]["value"] == "09171234567"      # Excel dropped the leading 0
    assert fields["email"]["normalized"] == "juan@example.com"
    assert fields["date_submitted"]["normalized"] == "2026-10-08"
    assert fields["contact_position"]["value"] == "Owner"
    assert not any(p["label"] == "Reference Type" for p in label_value_pairs(data))


def test_xlsx_detected_as_zip_is_accepted():
    check = check_file(as_plain_zip(cis_bytes()), "sheet.xlsx", 10**8)
    assert check.ok and check.extension == "xlsx"


def test_business_name_misread_is_likely_same():
    assert match_business("NISHIKEN GENERAL MERCHANDISE",
                          "NISHIKIKEN GENERAL MERCHANDISE") == LIKELY_SAME


def test_manila_epermit_valid_until_end_of_issue_year():
    v = check_validity("MAYORS_PERMIT", {"date_issued": "2026-02-04"}, TODAY)
    assert v.status == "VALID" and v.valid_until == "2026-12-31"
    assert v.issues[0].code == "VALIDITY_DERIVED"


def test_unknown_validity_of_required_document_blocks_active():
    p = permit()
    p.validity_status, p.valid_until = "UNKNOWN", None
    a = assess(SNAPSHOT, [bir(), dti(), p], TODAY)
    assert a["recommendation"] == "MANUAL_REVIEW"
    assert any("could not be determined: MAYORS_PERMIT" in r for r in a["reasons"])


def _sheet_doc(**overrides):
    sheet = extract_customer_info_sheet(cis_bytes(**overrides))
    return DocInput("cis", "cis.xlsx", sheet["document_type"], sheet["fields"], "NO_EXPIRY",
                    None, [], sheet["other_fields"])


def test_sheet_confirms_tin_and_fills_contacts():
    p = permit()
    p.other_fields = [{"label": "Telephone No.", "value": "91-6219-7943"},
                      {"label": "Email Address", "value": "shop@example.com"},
                      {"label": "No. of Emp.", "value": "3"}]
    a = assess({**SNAPSHOT, "Email": None, "Phone": None, "Contact_Person": None},
               [bir(), dti(), p, _sheet_doc()], TODAY)
    checks = {c["attribute"]: c for c in a["cross_document"]}
    assert checks["tin"]["status"] == "CONSISTENT"
    assert checks["email"]["status"] == "DIFFERS" and checks["email"]["severity"] == "INFO"
    assert a["recommendation"] == "ACTIVE"
    props = {p["zoho_field"]: p for p in a["proposals"]}
    assert props["Email"]["action"] == "FILL"
    assert props["Email"]["proposed_value"] == "juan@example.com"   # customer's own form first
    assert props["Phone"]["proposed_value"] == "09171234567"
    assert props["Contact_Person"]["proposed_value"] == "Juan Dela Cruz"
    data = {d["type"]: d for d in a["document_data"]}
    assert data["MAYORS_PERMIT"]["other_fields"][2] == {"label": "No. of Emp.", "value": "3"}
    assert data["CUSTOMER_INFO_SHEET"]["fields"]["tin"] == "123-456-789-00000"


def test_sheet_tin_different_from_bir_is_critical():
    a = assess(SNAPSHOT, [bir(), dti(), permit(), _sheet_doc(tin="999-999-999-00000")], TODAY)
    tin = next(c for c in a["cross_document"] if c["attribute"] == "tin")
    assert tin["status"] == "CONFLICT" and a["recommendation"] == "MANUAL_REVIEW"


def test_existing_zoho_contact_is_not_overwritten():
    a = assess({**SNAPSHOT, "Email": "old@example.com"}, [bir(), dti(), permit(),
                                                          _sheet_doc()], TODAY)
    email = next(p for p in a["proposals"] if p["zoho_field"] == "Email")
    assert email["action"] == "DIFFERS" and email["proposed_value"] is None


def test_contact_person_counts_as_owner_only_if_form_says_owner():
    sheet = _sheet_doc()
    sheet.fields["contact_position"]["value"] = "Purchasing"
    a = assess(SNAPSHOT, [bir(), dti(), permit(), sheet], TODAY)
    owners = [f_["document_type"] for f_ in a["facts"]["owner_name"]]
    assert "CUSTOMER_INFO_SHEET" not in owners


def test_try_extraction_reads_sheet_without_ai(tmp_path, capsys):
    from app.tools import try_extraction
    (tmp_path / "sheet.xlsx").write_bytes(cis_bytes())
    (tmp_path / ".env").write_text("EXTRACTION_PROVIDER=none\n", encoding="utf-8")
    try_extraction.main([str(tmp_path / "sheet.xlsx"), "--env", str(tmp_path / ".env")])
    out = capsys.readouterr().out
    assert "classified: CUSTOMER_INFO_SHEET" in out and "read from spreadsheet cells" in out


def test_pipeline_reads_and_extracts_sheet_without_ocr_or_ai(db, settings):
    from app.db.models import (
        Document,
        DocumentExtraction,
        DocumentSource,
        DocumentStatus,
        FileReading,
        JobStatus,
        StoredFile,
        VerificationJob,
    )
    from app.pipeline.extract import extract_job
    from app.pipeline.read import read_job
    from app.storage.base import LocalStorage, storage_key_for
    from tests.fake_vision import FakeVision

    storage = LocalStorage(settings.storage_local_dir)
    data = cis_bytes()
    sha = hashlib.sha256(data).hexdigest()
    storage.put(storage_key_for(sha, "xlsx"), data, "application/xlsx")
    db.add(StoredFile(sha256=sha, size_bytes=len(data), extension="xlsx",
                      mime_type="application/vnd.openxmlformats-officedocument.spreadsheetml"
                                ".sheet", storage_key=storage_key_for(sha, "xlsx")))
    job = VerificationJob(account_id="1000000000000001", status=JobStatus.COLLECTED)
    db.add(job)
    db.flush()
    db.add(Document(job_id=job.id, source=DocumentSource.ATTACHMENT, zoho_file_ref="1",
                    file_name="cis.xlsx", sha256=sha, status=DocumentStatus.STORED))
    db.commit()
    read_job(db, job, storage)
    assert db.query(FileReading).one().document_type == "CUSTOMER_INFO_SHEET"
    ai = FakeVision({}, "OTHER")
    extract_job(db, job, storage, ai, today=TODAY)
    row = db.query(DocumentExtraction).one()
    assert not ai.calls and row.model == "spreadsheet-cells"
    assert row.fields["tin"]["value"] == "123-456-789-00000"
